from __future__ import annotations
import csv
import io
from itertools import islice
import re
import zipfile
from openpyxl import load_workbook
from .schemas import address

MAX_ROWS=5000


def pasted(text):
    rows=[]
    for line in text.splitlines():
        if not line.strip(): continue
        match=re.fullmatch(r'\s*(.*?)\s*<([^<>]+)>\s*',line)
        if match: rows.append([match[1].strip(),match[2].strip()]); continue
        parts=re.split(r'[\t,，;；]+',line.strip())
        if len(parts)==1: parts=line.strip().rsplit(None,1)
        rows.append(parts)
        if len(rows)>MAX_ROWS+1: raise ValueError('一次最多导入 5000 行')
    return rows


def read_file(data,filename):
    if len(data)>5*1024*1024: raise ValueError('文件不能超过 5 MB')
    if filename.lower().endswith('.csv'):
        try: text=data.decode('utf-8-sig')
        except UnicodeDecodeError:
            try: text=data.decode('gb18030')
            except UnicodeDecodeError: raise ValueError('CSV 需使用 UTF-8 或 GB18030 编码') from None
        try: dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;\t')
        except csv.Error: dialect=csv.excel
        return list(islice(csv.reader(io.StringIO(text),dialect),MAX_ROWS+2))
    if filename.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist())>30*1024*1024: raise ValueError('XLSX 解压后体积过大')
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            try:
                if wb.active.max_column and wb.active.max_column>100: raise ValueError('文件列数过多')
                rows=[]
                for index,row in enumerate(wb.active.iter_rows(values_only=True)):
                    if index>MAX_ROWS: break
                    if len(row)>100: raise ValueError('文件列数过多')
                    rows.append(['' if c is None else str(c) for c in row])
                return rows
            finally: wb.close()
        except ValueError: raise
        except Exception: raise ValueError('无法读取 XLSX 文件，请确认文件未损坏且不是旧版 XLS') from None
    raise ValueError('只支持 CSV 或 XLSX 文件')


def preview(rows,existing,name_column=None,email_column=None,has_header=None):
    rows=[r for r in rows if any(str(v).strip() for v in r)]
    if len(rows)>MAX_ROWS+1: raise ValueError('一次最多导入 5000 行')
    if not rows: raise ValueError('没有找到名单，请粘贴内容或上传文件')
    first=[str(v).strip().lower() for v in rows[0]]
    names={'姓名','名字','学生姓名','name'}; emails={'邮箱','邮件','电子邮箱','email','e-mail'}
    ni=next((i for i,x in enumerate(first) if x in names),None)
    ei=next((i for i,x in enumerate(first) if x in emails),None)
    detected=ni is not None and ei is not None
    if name_column is None or email_column is None:
        if detected: name_column,email_column=ni,ei; has_header=True
        else:
            name_column,email_column=0,1
            if has_header is None: has_header=False
    if name_column==email_column or not 0<=name_column<100 or not 0<=email_column<100: raise ValueError('请选择两个不同的有效列')
    output=[]; seen={}; start=1 if has_header else 0
    for index,row in enumerate(rows[start:],start=start+1):
        name=str(row[name_column]).strip() if name_column<len(row) else ''
        email=str(row[email_column]).strip() if email_column<len(row) else ''
        status='valid'; message='可导入'
        if not name or len(name)>100 or '\n' in name or '\r' in name: status='invalid'; message='姓名为空、过长或包含换行'
        try: email=address(email)
        except ValueError: status='invalid'; message='邮箱格式不正确'
        if status=='valid':
            if email in seen:
                status='duplicate' if seen[email]==name else 'conflict'
                message='文件内重复邮箱，跳过' if status=='duplicate' else '文件内同邮箱姓名冲突，请修正源文件'
            elif email in existing:
                status='duplicate' if existing[email]==name else 'conflict'
                message='班级中已存在，跳过' if status=='duplicate' else '班级内同邮箱姓名不同，需确认覆盖'
            seen[email]=name
        output.append({'line':index,'name':name,'email':email,'status':status,'message':message})
    if len(output)>MAX_ROWS: raise ValueError('一次最多导入 5000 行')
    return {'rows':output,'columns':[str(x) for x in rows[0]],'has_header':bool(has_header),'name_column':name_column,'email_column':email_column,'auto_detected':detected}
